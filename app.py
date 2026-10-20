import streamlit as st
import google.generativeai as genai
import io, json
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils.cell import coordinate_to_tuple
from PIL import Image as PILImage
import pandas as pd
import xlsxwriter
from streamlit_paste_button import paste_image_button as pbutton

st.set_page_config(page_title="Préparation de commande", layout="centered")
st.title("Préparation de commande")

# --- Clé API Gemini ---
api_key = st.secrets.get("GEMINI_API_KEY", "")
if not api_key:
    api_key = st.sidebar.text_input("Clé API Gemini (Google)", type="password")

# --- Uploads ---
st.subheader("1. Fichiers")
template = st.file_uploader("Tableau Excel vierge (.xlsx)", type=["xlsx"])

st.write("**Photo de la commande** (capture d'écran)")
order_photo_result = pbutton("📋 Coller la photo de la commande", key="order_paste", errors="ignore")

st.write("**Photo du produit** (capture d'écran)")
product_photo_result = pbutton("📋 Coller la photo du produit", key="product_paste", errors="ignore")

st.subheader("2. Informations")
link = st.text_input("Lien Alibaba")
client_name = st.text_input("Nom du client (pour le nom du fichier)", placeholder="Ex. Simon Weber")

def extract_data_from_photo(image_bytes, api_key):
    genai.configure(api_key=api_key)
    model = genai.GenerativeModel('gemini-1.5-flash')
    img = PILImage.open(io.BytesIO(image_bytes))
    
    prompt = """
    Analyse cette image de commande (type Etsy). Extrait l'adresse de livraison globale, le nom du produit, et la LISTE des variantes (couleurs) avec leurs quantités.
    Renvoie UNIQUEMENT un JSON avec cette structure exacte :
    {
      "shipping": {
        "name": "Nom du destinataire",
        "street": "Numéro et rue",
        "address1": "Complément (Bâtiment, Casa...)",
        "city": "Ville",
        "zip": "Code postal",
        "province": "Province ou État",
        "country": "Pays"
      },
      "product_name": "Nom complet du produit (sans la couleur ni la quantité)",
      "variants": [
        {
          "color": "Couleur (ex: blanc, or, noir)",
          "quantity": 1
        }
      ]
    }
    Si l'image contient 4 produits avec 4 couleurs différentes, renvoie 4 objets dans la liste "variants".
    Si un produit n'a pas de couleur, mets "Standard".
    """
    response = model.generate_content([prompt, img])
    text = response.text.strip().replace("```json", "").replace("```", "").strip()
    return json.loads(text)

if st.button("Générer le fichier", type="primary"):
    if not all([template, link, client_name]):
        st.error("Veuillez remplir le tableau, le lien et le nom du client.")
        st.stop()
    
    if order_photo_result.image_data is None or product_photo_result.image_data is None:
        st.error("Veuillez coller les deux photos avec les boutons.")
        st.stop()

    # Conversion des images PIL en bytes
    order_img_bytes = io.BytesIO()
    order_photo_result.image_data.save(order_img_bytes, format="PNG")
    order_img_bytes = order_img_bytes.getvalue()

    product_img_bytes = io.BytesIO()
    product_photo_result.image_data.save(product_img_bytes, format="PNG")
    product_img_bytes = product_img_bytes.getvalue()

    with st.spinner("L'IA analyse la commande..."):
        data = extract_data_from_photo(order_img_bytes, api_key)
        shipping = data.get("shipping", {})
        product_name = data.get("product_name", "")
        variants = data.get("variants", [])

        if not variants:
            st.error("Aucune variante ou quantité détectée sur la photo.")
            st.stop()

    with st.spinner("Remplissage du tableau..."):
        wb = load_workbook(io.BytesIO(template.getvalue()))
        ws = wb.active
        
        start_row = 2
        
        # On boucle sur chaque variante
        for i, var in enumerate(variants):
            current_row = start_row + i
            
            # 1. Quantité et Couleur (toujours remplis)
            ws.cell(row=current_row, column=4, value=int(var.get("quantity", 1))) # D
            ws.cell(row=current_row, column=5, value=var.get("color", ""))        # E

            # 2. Les infos globales UNIQUEMENT sur la première ligne
            if i == 0:
                values = {
                    2: product_name,                     # B: Lineitem name
                    3: link,                             # C: Link Alibaba
                    6: shipping.get("name", ""),         # F: Shipping Name
                    7: shipping.get("street", ""),       # G: Shipping Street
                    8: shipping.get("address1", ""),     # H: Shipping Address1
                    10: shipping.get("city", ""),        # J: Shipping City
                    11: shipping.get("zip", ""),         # K: Shipping Zip
                    12: shipping.get("province", ""),    # L: Shipping Province
                    13: shipping.get("country", ""),     # M: Shipping Country
                }
                for col, val in values.items():
                    if val not in ("", None): 
                        ws.cell(row=current_row, column=col, value=val)

        # 3. Préparation de l'image du produit
        img = PILImage.open(io.BytesIO(product_img_bytes)).convert("RGB")
        img.thumbnail((100, 70)) 
        img_path = "temp_prod_img.png"
        img.save(img_path, format="PNG")

        # Sauvegarde avec openpyxl
        temp_out = io.BytesIO()
        wb.save(temp_out)
        
        # 4. Réécriture avec XlsxWriter pour l'image
        final_out = io.BytesIO()
        with pd.ExcelWriter(final_out, engine='xlsxwriter') as writer:
            df = pd.read_excel(io.BytesIO(temp_out.getvalue()), header=None)
            df.to_excel(writer, index=False, header=False)
            worksheet = writer.sheets['Sheet1']
            
            # Image sur la première ligne (P2)
            worksheet.insert_image(1, 15, img_path, {'x_scale': 0.8, 'y_scale': 0.8})
            
            # Hauteur pour toutes les lignes concernées
            for row_idx in range(1, 1 + len(variants)):
                worksheet.set_row(row_idx, 60)

        final_out.seek(0)

    st.success(f"Fichier généré avec {len(variants)} variante(s) !")
    st.download_button(
        "⬇️ Télécharger le fichier Excel", data=final_out,
        file_name=f"Cmde {client_name}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
